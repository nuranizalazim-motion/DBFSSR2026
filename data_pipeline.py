"""Read-only, traceable transformations for the supplied faculty sources."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import json
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pypdf import PdfReader

DIRECTORY = "FSSR DIREKTORI untuk Dr Nizam.xlsx"
PROGRAMMES = "Laporan_Program.xlsx"
GAME = "MAKLUMAT STAFF GAME + MOTION.xlsx"
STUDENTS = "pelajar Degree 2026.xlsx"
PDF = "SENARAI RUANG P&P KPSK PUNCAK ALAM (Makmal komputer dan Bilik Pembentangan).pdf"
PPT = "SENARAI BANGUNAN & RUANG MENGIKUT BLOK DI FSSR.pptx"
STAFF_FILES = {
    "MAKLUMATSTAF_IDE.xlsx": "IDE",
    "MAKLUMATSTAF_fesyen.xlsx": None,
    "MAKLUMATSTAF_JEWELLERY.xlsx": "Jewellery",
    "MAKLUMATSTAF_KOMUNIKASI VISUAL.xlsx": "Komunikasi Visual",
    "MAKLUMATSTAF_SERAMIK.xlsx": "Seramik",
}
REQUIRED = [DIRECTORY, PROGRAMMES, GAME, STUDENTS, *STAFF_FILES, PDF, PPT]
CAMPUS = {"A": "Seri Iskandar", "B": "Shah Alam", "B8": "Puncak Alam", "M": "Alor Gajah"}
MISSING = {"", "NA", "N/A", "NIL", "-"}


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clean(value):
    s = re.sub(r"\s+", " ", text(value))
    return None if s.upper() in MISSING else s


def number(value):
    if isinstance(value, bool) or clean(value) is None:
        return None
    try:
        n = float(text(value).replace(",", ""))
        return n if pd.notna(n) else None
    except ValueError:
        return None


def checked_number(value, context):
    n = number(value)
    if clean(value) is not None and n is None:
        raise ValueError(f"Unrecognised numeric value in {context}: {text(value)}. Refresh rejected; original retained.")
    return n


def name_key(value):
    # Only case, Unicode and whitespace. No fuzzy matching or title removal.
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text(value))).upper()


def year(value):
    if isinstance(value, (date, datetime)):
        return value.year
    s = text(value)
    if re.fullmatch(r"(?:19|20)\d{2}(?:\.0)?", s):
        return int(float(s))
    years = re.findall(r"\b(?:19|20)\d{2}\b", s)
    return int(years[0]) if len(set(years)) == 1 else None


def explicit_date(value):
    if isinstance(value, (date, datetime)):
        return pd.Timestamp(value).normalize()
    if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", text(value)):
        return pd.to_datetime(value, dayfirst=True, errors="coerce")
    return pd.NaT


def code_group(value):
    # Aliases in the SAME source cell, retained as a group; never explode counts.
    return "/".join(sorted(re.findall(r"(?:CAAD|AD)\d{3}", text(value).upper())))


def simple_sum(value):
    """Only explicit whole numbers/additions, never arbitrary expression evaluation."""
    s = text(value)
    if re.fullmatch(r"\d+(?:\s*\+\s*\d+)*", s):
        return sum(int(x) for x in re.findall(r"\d+", s))
    return None


def provenance(filename, sheet, row):
    return {"Source file": filename, "Source sheet/page": sheet, "Source row": row,
            "Record ID": f"{filename}:{sheet}:{row}"}


def person_block_end(sheet, row, name_column=3):
    for merged in sheet.merged_cells.ranges:
        if merged.min_row == row and merged.min_col <= name_column <= merged.max_col:
            return merged.max_row
    return row


def person_text(sheet, row, column):
    values = [clean(sheet.cell(i, column).value) for i in range(row, person_block_end(sheet, row) + 1)]
    return " | ".join(v for v in values if v) or None


def read_sources(folder: Path, replacements=None):
    replacements = replacements or {}
    unknown = set(replacements) - set(REQUIRED)
    if unknown:
        raise ValueError("Use the original source filenames. Unrecognised files: " + ", ".join(sorted(unknown)))
    result = {}
    for filename in REQUIRED:
        if filename in replacements:
            result[filename] = replacements[filename]
        elif (folder / filename).is_file():
            result[filename] = (folder / filename).read_bytes()
        else:
            raise ValueError(f"Missing source file: {filename}. Keep all 11 original filenames in the data folder.")
    return result


def profiles(sources):
    """Physical nonempty rows, not headcounts. No values from raw records in profile."""
    result = []
    for filename, content in sources.items():
        if not filename.endswith(".xlsx"):
            continue
        values = load_workbook(BytesIO(content), data_only=True)
        formulas = load_workbook(BytesIO(content), data_only=False)
        for s in values:
            header_rows = [3] if filename == DIRECTORY else [2, 3]
            if filename == STUDENTS:
                header_rows = [4, 5]
            elif filename == GAME:
                header_rows = [6, 7] if s.title != "JUMLAH PELAJAR" else [6]
            elif filename == PROGRAMMES:
                header_rows = [1, 2, 3] if s.title == "UNJUR DAFTAR 21-25" else ([2] if s.title == "I-ACCREDIT" else [3])
            data_rows = [r for r in s.iter_rows(min_row=max(header_rows) + 1)
                         if any(c.value is not None for c in r)]
            columns = []
            active_cols = sorted({c.column for r in s for c in r if c.value is not None})
            for j in active_cols:
                label = " / ".join(dict.fromkeys(text(s.cell(i, j).value) for i in header_rows if s.cell(i, j).value is not None))
                cells = [r[j - 1] for r in data_rows]
                filled = [c.value for c in cells if clean(c.value) is not None]
                columns.append({"column": get_column_letter(j), "header": label or "Unlabelled",
                                "missing": len(cells) - len(filled), "types": dict(Counter(type(v).__name__ for v in filled)),
                                "excel_date_formats": sorted({c.number_format for c in cells if getattr(c, "is_date", False)})})
            tuples = [tuple(text(c.value) for c in r) for r in data_rows]
            formula_cells = [c for row in formulas[s.title] for c in row if c.data_type == "f"]
            uncached = [c.coordinate for c in formula_cells if s[c.coordinate].value is None]
            result.append({"file": filename, "sheet": s.title, "physical_rows": s.max_row,
                           "nonempty_body_rows": len(data_rows), "exact_duplicate_body_rows": len(tuples) - len(set(tuples)),
                           "header_rows": header_rows, "columns": columns, "merged_ranges": [str(r) for r in s.merged_cells.ranges],
                           "formula_count": len(formula_cells), "uncached_formula_cells": uncached})
        values.close()
        formulas.close()
    return result


def checked_join(left, right, fields, label, audit):
    keys = ["Code group", "Campus"]
    if left.duplicated(keys).any() or right.duplicated(keys).any():
        raise ValueError(f"Unsafe join in {label}: duplicate programme + campus keys. Refresh rejected; correct source keys first.")
    missing = right.merge(left[keys], on=keys, how="left", indicator=True).query('_merge == "left_only"')
    if len(missing):
        raise ValueError(f"Unmatched programme + campus keys in {label}: {len(missing)}. Refresh rejected.")
    output = left.merge(right[keys + fields], on=keys, how="left", validate="one_to_one")
    audit.append({"Check": f"Join: {label}", "Outcome": "Passed", "Detail": f"{len(right)} unique keys; {len(left)} → {len(output)} base records. No fan-out."})
    return output


def build_snapshot(sources):
    missing = set(REQUIRED) - sources.keys()
    if missing:
        raise ValueError("Missing required sources: " + ", ".join(sorted(missing)))
    audit, staff, directory, facilities, students, wbl, history = [], [], [], [], [], [], []
    workbooks = {n: load_workbook(BytesIO(b), data_only=True) for n, b in sources.items() if n.endswith(".xlsx")}
    try:
        for filename, department in STAFF_FILES.items():
            for s in workbooks[filename]:
                if clean(s.cell(2, 3).value) != "NAMA":
                    raise ValueError(f"Unsupported staff header in {filename}/{s.title}")
                dept = department or s.title
                for i in range(4, s.max_row + 1):
                    r = [s.cell(i, j).value for j in range(1, 11)]
                    if number(r[1]) is None or clean(r[2]) is None:
                        continue
                    staff.append({"Name": clean(r[2]), "Department": dept, "Employee ID": None,
                                  "Grade": clean(r[3]), "Qualifications (source)": person_text(s, i, 5),
                                  "Study registration (source)": person_text(s, i, 6), "Expected completion (source)": person_text(s, i, 7),
                                  "Funding (source)": person_text(s, i, 8), "Age (source snapshot)": checked_number(r[8], f"{filename}/{i} age"),
                                  "Retirement (source)": clean(r[9]), "Retirement year": year(r[9]),
                                  **provenance(filename, s.title, i)})
        for sn in ["GAME DESIGN", "MOTION DESIGN"]:
            s = workbooks[GAME][sn]
            if clean(s.cell(6, 3).value) != "NAMA":
                raise ValueError(f"Unsupported staff header in {GAME}/{sn}")
            for i in range(8, s.max_row + 1):
                r = [s.cell(i, j).value for j in range(1, 14)]
                if number(r[0]) is None or clean(r[2]) is None:
                    continue
                # Use this person's actual merged name range, never a guessed block length.
                staff.append({"Name": clean(r[2]), "Department": sn.title(), "Employee ID": clean(r[1]),
                              "Grade": clean(r[4]), "Qualifications (source)": person_text(s, i, 6),
                              "Study registration (source)": person_text(s, i, 8), "Expected completion (source)": person_text(s, i, 9),
                              "Funding (source)": person_text(s, i, 10), "Age (source snapshot)": checked_number(r[10], f"{GAME}/{sn}/{i} age"),
                              "Retirement (source)": clean(r[11]), "Retirement year": year(r[11]),
                              **provenance(GAME, sn, i)})
        sf = pd.DataFrame(staff)
        if sf.empty:
            raise ValueError("No staff records found; check the source templates.")
        sf["Name key"] = sf["Name"].map(name_key)
        sf["Possible duplicate name"] = sf["Name key"].duplicated(keep=False)
        sf["Qualification label"] = sf["Qualifications (source)"].map(lambda v: "Missing" if v is None else
            "PhD mentioned" if re.search(r"P\s*H\.?\s*D|DOCTOR OF PHILOSOPHY|PHILOSOPHIAE DOCTOR", v, re.I) else
            "Master mentioned" if re.search(r"MASTER|\bMA\b", v, re.I) else "Other / review source")
        sf["Quality note"] = sf.apply(lambda r: "; ".join(filter(None, [
            "Possible duplicate name; identity not merged" if r["Possible duplicate name"] else "",
            "Unusual grade retained" if r["Grade"] not in ("DS11", "DS13", "DS14", "VK", None) else "",
            "Retirement year ≥2060: verify source" if pd.notna(r["Retirement year"]) and r["Retirement year"] >= 2060 else "",
            "Qualifications missing" if not r["Qualifications (source)"] else ""])), axis=1)
        for s in workbooks[DIRECTORY]:
            section = "Unspecified"
            for i, row in enumerate(s.values, 1):
                r = list(row)
                if clean(r[0]) and not clean(r[1]):
                    section = clean(r[0])
                elif number(r[0]) is not None and clean(r[1]):
                    directory.append({"Name": clean(r[1]), "Role": clean(r[2]), "Phone (source)": clean(r[3]),
                                      "Section": section, "Staff category": "Academic" if s.title.startswith("AKADEMIK") else "Administration",
                                      **provenance(DIRECTORY, s.title, i)})
        df = pd.DataFrame(directory)
        df["Possible duplicate name"] = df["Name"].map(name_key).duplicated(keep=False)
        ps = workbooks[PROGRAMMES]
        def programme_rows(sn, start):
            output = []
            for i in range(start, ps[sn].max_row + 1):
                r = [c.value for c in ps[sn][i]]
                if number(r[0]) is not None:
                    if not any(clean(v) for v in r[1:]):
                        audit.append({"Check": f"Blank template row: {sn}/{i}", "Outcome": "Excluded", "Detail": "Numbered row with no programme data; not counted as a programme."})
                        continue
                    if not code_group(r[1]) or not clean(r[3]):
                        raise ValueError(f"Missing programme or campus key in {sn}, row {i}. Refresh rejected.")
                    output.append((i, r))
            return output
        base = []
        for i, r in programme_rows("I-ACCREDIT", 3):
            item = {"Code group": code_group(r[1]), "Programme": clean(r[2]), "Campus": clean(r[3]),
                    "Accreditation reference": clean(r[6]), **provenance(PROGRAMMES, "I-ACCREDIT", i)}
            for label, j in [("Programme start", 7), ("Accreditation start", 8), ("Accreditation end", 9), ("Maturity date", 11)]:
                item[label] = explicit_date(r[j]); item[label + " (source)"] = clean(r[j])
            base.append(item)
        pf = pd.DataFrame(base)
        if pf.empty or pf.duplicated(["Code group", "Campus"]).any():
            raise ValueError("Programme base has empty or duplicate keys.")
        counts = pd.DataFrame([{ "Code group": code_group(r[1]), "Campus": clean(r[3]),
            "Senior staff (source)": checked_number(r[4], f"programme row {i} senior count"), "DM11 staff (source)": checked_number(r[5], f"programme row {i} DM11 count"), "Students (report source)": checked_number(r[6], f"programme row {i} student count"),
            "Staff count row": i} for i, r in programme_rows("Maklumat staf x student", 4)])
        pf = checked_join(pf, counts, [c for c in counts if c not in ["Code group", "Campus"]], "Maklumat staf x student", audit)
        review = pd.DataFrame([{ "Code group": code_group(r[1]), "Campus": clean(r[3]), "Review stage": clean(r[4]),
            "Review status": clean(r[5]), "Review date": explicit_date(r[8]), "Review date (source)": clean(r[8]), "Review row": i}
            for i, r in programme_rows("Maklumat Semakan", 4)])
        pf = checked_join(pf, review, [c for c in review if c not in ["Code group", "Campus"]], "Maklumat Semakan", audit)
        consolidation = pd.DataFrame([{ "Code group": code_group(r[1]), "Campus": clean(r[3]), "Consolidation direction": clean(r[4]),
            "Consolidation status": clean(r[5]), "Committee chair": clean(r[6]), "Last meeting (source)": clean(r[7]), "Consolidation row": i}
            for i, r in programme_rows("Konsolidisasi_WBL", 4)])
        pf = checked_join(pf, consolidation, [c for c in consolidation if c not in ["Code group", "Campus"]], "Konsolidisasi_WBL", audit)
        pf["Staff total (complete pairs only)"] = pf[["Senior staff (source)", "DM11 staff (source)"]].sum(axis=1, min_count=2)
        pf["Students per reported staff"] = pf["Students (report source)"] / pf["Staff total (complete pairs only)"].where(pf["Staff total (complete pairs only)"] > 0)
        pf["Review status"] = pf["Review status"].fillna("Not supplied")
        s = workbooks[STUDENTS].active
        if text(s.cell(5, 4).value) != "A" or text(s.cell(5, 7).value) != "M":
            raise ValueError("Student campus headers changed. Refresh rejected.")
        for i in range(6, s.max_row + 1):
            if not code_group(s.cell(i, 3).value):
                continue
            name = clean(s.cell(i, 3).value)
            total = checked_number(s.cell(i, 8).value, f"degree row {i} total")
            subtotal = 0
            for j, campus in enumerate(["A", "B", "B8", "M"], 4):
                n = checked_number(s.cell(i, j).value, f"degree row {i} campus {campus}")
                if n is not None:
                    subtotal += n
                students.append({"Programme": name, "Campus code": campus, "Campus": CAMPUS[campus],
                                 "Students": n, "Source programme total (not additive)": total, **provenance(STUDENTS, s.title, f"{i}:{get_column_letter(j)}")})
            audit.append({"Check": f"Degree programme row {i}", "Outcome": "Passed" if total == subtotal else "Mismatch",
                          "Detail": f"Nonblank campuses sum {subtotal:g}; source total {total}. Blanks remain missing."})
        ef = pd.DataFrame(students)
        totalrows = [i for i in range(1, s.max_row + 1) if text(s.cell(i, 3).value).upper() == "JUMLAH KESELURUHAN"]
        if len(totalrows) != 1:
            raise ValueError("Expected one degree grand-total row. Refresh rejected.")
        totalrow = totalrows[0]
        source_total = checked_number(s.cell(totalrow, 8).value, "degree grand total")
        audit.append({"Check": "Degree grand total", "Outcome": "Passed" if ef.Students.sum() == source_total else "Mismatch",
                      "Detail": f"Calculated {ef.Students.sum():g}; source {source_total}. Excludes subtotal/grand-total rows."})
        for j, campus in enumerate(["A", "B", "B8", "M"], 4):
            actual = ef.loc[ef["Campus code"] == campus, "Students"].sum()
            reported = checked_number(s.cell(totalrow, j).value, f"degree campus {campus} grand total")
            audit.append({"Check": f"Degree campus {campus}", "Outcome": "Passed" if actual == reported else "Mismatch",
                          "Detail": f"Calculated {actual:g}; source {reported}."})
        s = workbooks[GAME]["JUMLAH PELAJAR"]
        for offset, programme, totalrow in [(1, "Game Design", 15), (8, "Motion Design", 18)]:
            semester = None
            for i in range(7, totalrow):
                rawsem = clean(s.cell(i, offset + 1).value)
                if rawsem and rawsem.startswith("SEM"):
                    semester = rawsem
                group = clean(s.cell(i, offset + 3).value)
                if not group or not re.fullmatch(r"(?:CAAD|AD)\d+[A-Z]", group):
                    continue
                raw = clean(s.cell(i, offset + 2).value)
                parsed = simple_sum(raw)
                wbl.append({"Programme": programme, "Session": "20262 (Mac–Ogos 2026)", "Semester (source)": semester,
                            "Group": group, "Students": parsed, "Count (source)": raw,
                            "Advisor (source)": clean(s.cell(i, offset + 6).value),
                            **provenance(GAME, s.title, f"{i}:{offset + 2}")})
            rows = [r for r in wbl if r["Programme"] == programme]
            parsedtotal = sum(r["Students"] for r in rows if r["Students"] is not None)
            stated = number(s.cell(totalrow, offset + 2).value)
            audit.append({"Check": f"WBL {programme} source total", "Outcome": "Review" if any(r["Students"] is None for r in rows) else ("Passed" if parsedtotal == stated else "Mismatch"),
                          "Detail": f"Unambiguous group sum {parsedtotal}; source total {stated:g}. Annotated counts excluded; do not add to degree totals."})
        # Retain both historical grains; dashboard requires one grain at a time.
        hs = ps["UNJUR DAFTAR 21-25"]
        starts = [(c.column, re.search(r"\((\d{5})\)", text(c.value)).group(1)) for c in hs[2]
                  if re.search(r"\((\d{5})\)", text(c.value))]
        groups, current = [], None
        for i in range(4, hs.max_row + 1):
            r = [c.value for c in hs[i]]
            if code_group(r[1]):
                current = {"name": clean(r[2]), "code": code_group(r[1]), "parent": (i, r), "children": []}
                groups.append(current)
            elif current and text(r[2]).upper().startswith("UITM"):
                current["children"].append((i, r))
        for group in groups:
            for i, r in [group["parent"], *group["children"]]:
                is_parent = i == group["parent"][0]
                campus = "All campuses (source total)" if is_parent else clean(r[2])
                for k, (col, session) in enumerate(starts):
                    end = starts[k + 1][0] if k + 1 < len(starts) else hs.max_column + 1
                    values = {text(hs.cell(3, j).value): r[j - 1] for j in range(col, end)}
                    projected, registered, offered = (number(values.get(key)) for key in ["UNJ", "DFT", "TWR"])
                    if all(v is None for v in [projected, registered, offered]):
                        continue
                    history.append({"Programme": group["name"], "Code group": group["code"], "Campus": campus,
                        "Record level": "Programme totals" if is_parent else "Campus breakdown",
                        "Session": session, "Year": int(session[:4]), "Period": f"{session[:4]}-{'03' if session.endswith('2') else '10'}",
                        "Projected (UNJ)": projected, "Registered (DFT)": registered, "Offered (TWR)": offered,
                        "Source variance %": number(values.get("% DFT vs UNJ")), **provenance(PROGRAMMES, hs.title, i)})
        hf = pd.DataFrame(history)
        breakdown_mismatches = 0
        for _, parent in hf[hf["Record level"] == "Programme totals"].iterrows():
            children = hf[(hf["Record level"] == "Campus breakdown") & (hf["Code group"] == parent["Code group"]) & (hf.Session == parent.Session)]
            for column in ["Projected (UNJ)", "Registered (DFT)", "Offered (TWR)"]:
                known = children[column].dropna()
                if len(known) and pd.notna(parent[column]) and known.sum() != parent[column]:
                    breakdown_mismatches += 1
        audit.append({"Check": "Historical campus reconciliation", "Outcome": "Review" if breakdown_mismatches else "Passed",
                      "Detail": f"{breakdown_mismatches} parent metric cells differ from known campus sums. Programme totals and campus breakdowns are separate selectable grains, never combined."})
        hf["Variance % (calculated)"] = 100 * (hf["Registered (DFT)"] - hf["Projected (UNJ)"]) / hf["Projected (UNJ)"].where(hf["Projected (UNJ)"] != 0)
        comparable = hf["Variance % (calculated)"].notna() & hf["Source variance %"].notna()
        disagreements = (abs(hf.loc[comparable, "Variance % (calculated)"] - hf.loc[comparable, "Source variance %"]) > .1).sum()
        audit.append({"Check": "Historical source percentages", "Outcome": "Mismatch" if disagreements else "Passed",
                      "Detail": f"{int(comparable.sum())} comparable cells; {disagreements} differ by >0.1 percentage point. Displayed variance recalculated; raw retained."})
        audit.append({"Check": "Staff identity", "Outcome": "Review", "Detail": f"{len(sf)} records, {sf['Name key'].nunique()} exact-normalised names; {int(sf['Possible duplicate name'].sum())} rows flagged. No confirmed unique-person total; {sf['Employee ID'].notna().sum()} source employee IDs."})
        audit.append({"Check": "Directory identity", "Outcome": "Review", "Detail": f"{len(df)} directory entries; {int(df['Possible duplicate name'].sum())} repeated-name entries; kept separate from staff planning."})
        audit.append({"Check": "Coverage across student sources", "Outcome": "Review", "Detail": "Degree, WBL and programme-report counts are distinct snapshots with different coverage. No shared source dates establish equivalence; no cross-source sums or staffing ratios. Compare labelled views, not combined totals."})
        audit.append({"Check": "Excluded worksheet", "Outcome": "Excluded", "Detail": "Laporan_Program/test repeats accreditation information. Profiled but excluded from analytical totals."})
    finally:
        for workbook in workbooks.values():
            workbook.close()
    # PDF text contains explicit room IDs and count/capacity/floor columns.
    reader = PdfReader(BytesIO(sources[PDF]))
    for page_index, page in enumerate(reader.pages, 1):
        raw = page.extract_text() or ""
        chunks = re.split(r"(?m)(?=^\d+\s+\d+\s+A\s+\d+\s+\d+/0\b)", raw)
        page_count = 0
        for chunk in chunks:
            line = re.sub(r"\s+", " ", chunk).strip()
            if "KOLEJ PENGAJIAN" in line:
                line = line.split("KOLEJ PENGAJIAN")[0].strip()
            m = re.fullmatch(r"(\d+)\s+(\d+\s+A\s+\d+\s+\d+/0)\s+(.+?)\s+(\d+)\s+(\d+)", line)
            if not m:
                if re.match(r"\d+\s+\d+\s+A\s+\d+\s+\d+/0", line):
                    raise ValueError("A PDF room row could not be parsed. Refresh rejected rather than silently losing it.")
                continue
            seq, identifier, desc, capacity, floor = m.groups()
            equipment = re.search(r"\s(\d+)$", desc) if re.search(r"iMac|\bPC\b|DEKSTOP|ALIANWARE", desc, re.I) else None
            facilities.append({"Campus": "Puncak Alam", "Block": "A", "Floor (source)": floor,
                "Room / group": desc, "Room ID": identifier, "Rooms": 1, "Student capacity (source)": int(capacity),
                "Equipment units (source)": int(equipment.group(1)) if equipment else None,
                "Other capacity (source)": None, "Capacity text (source)": capacity, "Quality note": "",
                **provenance(PDF, f"Page {page_index}", int(seq))})
            page_count += 1
        if not page_count:
            raise ValueError("PDF contains no extractable room rows; OCR/template review needed.")
    pr = Presentation(BytesIO(sources[PPT]))
    for idx, slide in enumerate(pr.slides, 1):
        titles = " ".join(sh.text for sh in slide.shapes if sh.has_text_frame)
        block = re.search(r"BILANGAN RUANG DI BLOK\s+([A-H])", titles)
        if not block:
            continue
        floor = "Not supplied"
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            for row_number, row in enumerate(list(shape.table.rows)[1:], 2):
                cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells]
                if len(cells) != 3 or not cells[1] or not cells[2]:
                    continue
                if cells[0].startswith("ARAS"):
                    floor = cells[0]
                leading = cells[2].split("(")[0].strip()
                if "=" in leading:
                    expression, declared = leading.split("=", 1)
                    rooms = simple_sum(expression)
                    if rooms != simple_sum(declared):
                        raise ValueError(f"Room arithmetic disagreement in slide {idx}")
                else:
                    rooms = simple_sum(leading)
                if rooms is None:
                    raise ValueError(f"Unsupported room count in slide {idx}: {leading}")
                student_capacity = re.search(r"\((\d+)\s+Pelajar\)", cells[2], re.I)
                facilities.append({"Campus": "Shah Alam", "Block": block.group(1), "Floor (source)": floor,
                    "Room / group": cells[1], "Room ID": None, "Rooms": rooms,
                    "Student capacity (source)": int(student_capacity.group(1)) if student_capacity else None,
                    "Equipment units (source)": None, "Other capacity (source)": None if student_capacity else cells[2],
                    "Capacity text (source)": cells[2], "Quality note": "Source says function changed; not updated" if "BELUM KEMASKINI" in cells[1] else "",
                    **provenance(PPT, f"Slide {idx}", row_number)})
    ff = pd.DataFrame(facilities)
    pdfrows = ff[ff["Source file"] == PDF]
    if pdfrows["Room ID"].duplicated().any():
        raise ValueError("Duplicate PDF room IDs. Refresh rejected to prevent double counting.")
    audit.append({"Check": "Facilities extraction", "Outcome": "Passed", "Detail": f"{len(pdfrows)} PDF rooms; {len(ff)-len(pdfrows)} PowerPoint groups from {len(pr.slides)} inspected slides. Group capacities not multiplied by room counts; staff capacities excluded from student capacities."})
    profile = profiles(sources)
    for p in profile:
        if p["uncached_formula_cells"]:
            audit.append({"Check": f"Uncached formulas: {p['file']}/{p['sheet']}", "Outcome": "Review",
                          "Detail": f"{len(p['uncached_formula_cells'])} formula cells have no cached value; treated as missing, not evaluated."})
    for name, frame in [("Staff", sf), ("Directory", df), ("Programmes", pf), ("Students", ef), ("WBL", pd.DataFrame(wbl)), ("History", hf), ("Facilities", ff)]:
        if frame.empty:
            raise ValueError(f"No {name.lower()} records parsed. Refresh rejected.")
    now = datetime.now(timezone.utc).isoformat()
    return {"tables": {"Students": ef, "Programmes": pf, "Staff planning": sf, "Directory": df,
                       "WBL groups": pd.DataFrame(wbl), "Historical intake": hf, "Facilities": ff},
            "audit": pd.DataFrame(audit), "profiles": profile, "loaded_at": now,
            "manifest": [{"File": n, "Bytes": len(b), "SHA256": sha256(b).hexdigest()} for n, b in sources.items()]}


def refresh_snapshot(state, loader):
    """Atomic in-memory refresh: failed attempts preserve the last validated snapshot."""
    state["last_attempt"] = datetime.now(timezone.utc).isoformat()
    try:
        candidate = loader()
    except Exception as exc:
        state["refresh_error"] = str(exc)
        return False
    state["snapshot"] = candidate
    state["refresh_error"] = None
    return True


def apply_filters(frame, categories=None, query="", date_filter=None, number_filter=None):
    result = frame.copy()
    for col, selected in (categories or {}).items():
        # Empty selection intentionally yields no records; never broadens to all.
        result = result[result[col].fillna("Not supplied").astype(str).isin(selected)]
    if query.strip():
        searchable = result.astype(str).apply(lambda c: c.str.contains(query.strip(), case=False, regex=False))
        result = result[searchable.any(axis=1)]
    if date_filter:
        col, first, last = date_filter
        parsed = pd.to_datetime(result[col], errors="coerce")
        result = result[parsed.between(pd.Timestamp(first), pd.Timestamp(last))]
    if number_filter:
        col, first, last = number_filter
        numeric = pd.to_numeric(result[col], errors="coerce")
        result = result[numeric.between(first, last)]
    return result.reset_index(drop=True)


def safe_export(frame):
    result = frame.copy()
    for col in result.select_dtypes(include=["object", "string"]):
        result[col] = result[col].map(lambda v: "'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else v)
    return result


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, default=Path(__file__).parent / "data")
    parser.add_argument("--output", type=Path, default=Path(__file__).parent / "reports")
    args = parser.parse_args()
    snapshot = build_snapshot(read_sources(args.data_dir))
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "source_profile.json").write_text(json.dumps(snapshot["profiles"], indent=2), encoding="utf-8")
    (args.output / "manifest.json").write_text(json.dumps(snapshot["manifest"], indent=2), encoding="utf-8")
    snapshot["audit"].to_csv(args.output / "validation.csv", index=False)
    for name, table in snapshot["tables"].items():
        safe_export(table).to_csv(args.output / (name.lower().replace(" ", "_") + ".csv"), index=False)
    print(json.dumps({"rows": {n: len(f) for n, f in snapshot["tables"].items()},
                      "degree_students": float(snapshot["tables"]["Students"].Students.sum()),
                      "checks": snapshot["audit"].Outcome.value_counts().to_dict()}, indent=2))
